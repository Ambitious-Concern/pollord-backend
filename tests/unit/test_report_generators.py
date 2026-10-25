import codecs
import io

from openpyxl import load_workbook

from app.services.report_export_service import _derive_key
from app.schemas.report_export import ReportExportRequest
from app.utils.report_csv_generator import build_csv_bytes
from app.utils.report_document import MONEY, ReportDocument, ReportSection
from app.utils.report_excel_generator import build_excel_bytes
from app.utils.report_pdf_generator import build_pdf_bytes


def _doc(*sections) -> ReportDocument:
    return ReportDocument(
        title="Test", subtitle="sub", generated_at="now", sections=list(sections)
    )


class TestCsvGenerator:
    def test_produces_parseable_csv_with_bom(self):
        content = build_csv_bytes(
            _doc(
                ReportSection(
                    heading="Nominees",
                    headers=["Name", "Votes"],
                    rows=[["Ama", 10], ["Kofi", 7]],
                )
            )
        )
        assert content.startswith(codecs.BOM_UTF8)
        text = content.decode("utf-8-sig")
        assert "Name,Votes" in text
        assert "Ama,10" in text

    def test_includes_every_section_not_just_the_largest(self):
        """Regression: CSV used to emit only the single most detailed table,
        silently dropping the summary/breakdown tables that the Excel and PDF
        versions of the same report both include."""
        content = build_csv_bytes(
            _doc(
                ReportSection(heading="Summary", headers=["Metric", "Value"], rows=[["Gross", 100]]),
                ReportSection(heading="By channel", headers=["Channel", "Amount"], rows=[["card", 60]]),
                ReportSection(heading="Transactions", headers=["Ref"], rows=[["abc"], ["def"]]),
            )
        )
        text = content.decode("utf-8-sig")
        for heading in ("Summary", "By channel", "Transactions"):
            assert heading in text
        assert "Gross,100" in text
        assert "card,60" in text
        assert "abc" in text and "def" in text


class TestExcelGenerator:
    def test_one_sheet_per_section(self):
        document = _doc(
            ReportSection(heading="Summary", headers=["Metric", "Value"], rows=[["Gross", 100]]),
            ReportSection(heading="Transactions", headers=["Ref"], rows=[["abc"], ["def"]]),
        )
        workbook = load_workbook(io.BytesIO(build_excel_bytes(document)))
        assert workbook.sheetnames == ["Summary", "Transactions"]
        # A title block occupies the first rows, so the table starts lower.
        assert workbook["Summary"]["A5"].value == "Metric"
        assert workbook["Transactions"]["A6"].value == "abc"
        assert "Test" in str(workbook["Summary"]["A1"].value)

    def test_numeric_columns_stay_numeric_with_a_format(self):
        """A sheet of money-shaped *text* can't be summed, sorted or filtered,
        which defeats the point of exporting to a spreadsheet at all."""
        document = _doc(
            ReportSection(
                heading="Transactions",
                headers=["Ref", "Gross"],
                rows=[["a", 25.5], ["b", 10.0]],
                number_formats={1: MONEY},
                total_columns=[1],
            )
        )
        sheet = load_workbook(io.BytesIO(build_excel_bytes(document)))["Transactions"]
        assert sheet["B6"].value == 25.5
        assert isinstance(sheet["B6"].value, float)
        assert sheet["B6"].number_format == MONEY
        # A live formula, so the total still holds when rows are filtered.
        assert str(sheet["B8"].value).startswith("=SUBTOTAL(109")

    def test_sanitizes_unsafe_sheet_names(self):
        document = _doc(ReportSection(heading="Bad:Name/Here", headers=["X"], rows=[]))
        workbook = load_workbook(io.BytesIO(build_excel_bytes(document)))
        assert workbook.sheetnames == ["BadNameHere"]


class TestPdfGenerator:
    def test_produces_valid_pdf_bytes(self):
        document = ReportDocument(
            title="Voting Results", subtitle="Test Election", generated_at="2026-09-17",
            sections=[ReportSection(heading="Best Dressed", headers=["Rank", "Nominee"], rows=[[1, "Ama"]])],
        )
        content = build_pdf_bytes(document)
        assert content.startswith(b"%PDF")
        assert len(content) > 500


class TestIdempotencyKeyDerivation:
    def test_same_params_same_key(self):
        req = ReportExportRequest(report_type="financial", format="pdf", params={"range": "all-time"})
        assert _derive_key(req) == _derive_key(req)

    def test_different_params_different_key(self):
        req1 = ReportExportRequest(report_type="financial", format="pdf", params={"range": "all-time"})
        req2 = ReportExportRequest(report_type="financial", format="pdf", params={"range": "last-30"})
        assert _derive_key(req1) != _derive_key(req2)

    def test_different_format_different_key(self):
        req1 = ReportExportRequest(report_type="financial", format="pdf", params={})
        req2 = ReportExportRequest(report_type="financial", format="csv", params={})
        assert _derive_key(req1) != _derive_key(req2)
